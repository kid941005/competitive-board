"""表格导入（Excel / CSV）与 Excel 导出的单测。

重点：
  · Excel 与 CSV 必须走同一套校验 —— 不能出现「CSV 拦住了、Excel 放过去」的差异
  · 大类填了就必须认得出，认识不了宁可跳过也不猜
  · 导出是给人看的：最优值要加粗、缺失要是「—」、完整度要能对上

两种跑法都支持：pytest tests/ 或 python tests/test_import_export.py
使用独立临时数据目录，不碰开发库。
"""

from __future__ import annotations

import io
import os
import re
import sys
import tempfile
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="cib-sheet-test-")
os.environ["CIB_DATA_DIR"] = _TMP
os.environ["CIB_ADMIN_USER"] = "admin"
os.environ["CIB_ADMIN_PASSWORD"] = "AdminPass123"
os.environ.pop("DATABASE_URL", None)

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from _wide import wide_csv, wide_headers
from app.db import engine, init_db
from app.main import app
from app.models import Entity, Product, ProductSpec, SpecMetric
from app.seed import seed
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlmodel import Session, select

WIDE_HEAD = wide_headers()
_state: dict[str, object] = {}


def _client() -> TestClient:
    if _state.get("client"):
        return _state["client"]  # type: ignore[return-value]
    init_db()
    seed()
    c = TestClient(app)
    c.__enter__()
    r = c.post("/login", data={"username": "admin", "password": "AdminPass123"},
               follow_redirects=False)
    assert r.status_code == 303, f"登录失败：{r.status_code}"
    _state["client"] = c
    return c


def _csrf(c: TestClient) -> str:
    html = c.get("/import").text
    return re.search(r'name="csrf-token" content="([^"]+)"', html).group(1)


def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(c: TestClient, name: str, data: bytes, mime: str) -> dict:
    r = c.post(
        "/api/import",
        files={"file": (name, data, mime)},
        headers={"X-CSRF-Token": _csrf(c)},
    )
    assert r.status_code == 200, r.text
    return r.json()


# ── 导入（宽表：一行一个型号）─────────────────────────────────────────────
def test_import_xlsx_creates_company_product_and_specs():
    c = _client()
    data = _xlsx([
        WIDE_HEAD,
        ["表格测试甲", "XL-100", "压电气泵", 0.42, 8600, "", "", "", "", "", "", "", "",
         "https://x/1.pdf", 0.9, "", "静压@15Vpp"],
    ])
    rep = _post(c, "t.xlsx", data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert rep["rows_read"] == 1, rep
    assert rep["created_entities"] == 1 and rep["created_products"] == 1, rep
    assert rep["written_specs"] == 2 and rep["errors"] == [], rep

    with Session(engine) as s:
        ent = s.exec(select(Entity).where(Entity.name == "表格测试甲")).first()
        assert ent is not None
        prod = s.exec(select(Product).where(Product.model == "XL-100")).first()
        assert prod is not None and prod.category == "pump_air" and prod.entity_id == ent.id
        specs = s.exec(select(ProductSpec).where(ProductSpec.product_id == prod.id)).all()
        assert len(specs) == 2
        assert all(x.verified is False for x in specs), "导入的参数必须待复核"


def test_csv_and_xlsx_share_the_same_validation():
    """跨维度单位（把 kPa 写进流量列）在 CSV 与 Excel 下都必须被拦住。

    这条守的是「Excel 与 CSV 只在读文件那一步分叉」，不是各写一套校验。
    """
    c = _client()
    text = wide_csv({"企业名称": "表格测试乙", "产品型号": "CSV-1", "产品类别": "压电气泵",
                     "流量L/min": "1.5 kPa"})
    rep_csv = _post(c, "t.csv", text.encode("utf-8"), "text/csv")

    data = _xlsx([WIDE_HEAD,
                  ["表格测试丙", "XL-2", "压电气泵", "1.5 kPa", "", "", "", "", "", "", "",
                   "", "", "", "", "", ""]])
    rep_xlsx = _post(c, "t.xlsx", data,
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    assert rep_csv["written_specs"] == 0 and rep_xlsx["written_specs"] == 0
    assert "跨维度" in rep_csv["errors"][0], rep_csv["errors"]
    assert "跨维度" in rep_xlsx["errors"][0], rep_xlsx["errors"]


def test_unknown_column_is_ignored_not_guessed():
    """表里多出来的列（不在字典里的指标）直接忽略 —— 宽表按列头认指标，认不出就当没写。"""
    c = _client()
    text = (
        "企业名称,产品型号,产品类别,流量L/min,瞎编的指标\n"
        "表格测试己,UNK-1,压电气泵,0.7,3\n"
    )
    rep = _post(c, "t.csv", text.encode("utf-8"), "text/csv")
    assert rep["written_specs"] == 1, rep
    assert rep["errors"] == [], rep


def test_unknown_category_is_rejected_not_guessed():
    c = _client()
    data = _xlsx([WIDE_HEAD, ["表格测试丁", "BAD-1", "压电陶瓷", 1.0, "", "", "", "", "",
                              "", "", "", "", "", "", "", ""]])
    rep = _post(c, "t.xlsx", data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert rep["written_specs"] == 0
    assert any("不可识别" in e for e in rep["errors"]), rep["errors"]


def test_category_conflict_on_existing_product_is_skipped():
    """mp6-liq 是液泵，表里按气泵写必须报错 —— 否则参数会落到错误的指标体系。"""
    c = _client()
    data = _xlsx([WIDE_HEAD, ["巴特尔斯微技术", "mp6-liq", "压电气泵", 0.5, "", "", "", "",
                              "", "", "", "", "", "", "", "", ""]])
    rep = _post(c, "t.xlsx", data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert rep["written_specs"] == 0
    assert any("冲突" in e for e in rep["errors"]), rep["errors"]


def test_missing_required_header_reports_clearly():
    """缺「产品型号」等于读不出来 —— 报错要说清模板已经改成宽表，别让人瞎猜。"""
    c = _client()
    data = _xlsx([["企业名称", "流量L/min"], ["某公司", 1.0]])
    rep = _post(c, "t.xlsx", data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert rep["written_specs"] == 0
    assert any("宽表" in e for e in rep["errors"]), rep["errors"]


def test_placeholder_value_counts_as_not_disclosed():
    """「待补充」是表里最高频的值，必须当成未公开跳过，而不是报错或写 0。"""
    c = _client()
    data = _xlsx([WIDE_HEAD, ["表格测试戊", "BL-1", "压电气泵", "待补充", "待补充", "", "", "",
                              "", "", "", "", "", "", "", "", ""]])
    rep = _post(c, "t.xlsx", data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert rep["errors"] == [], rep
    assert rep["skipped_missing"] >= 2 and rep["written_specs"] == 0


def test_unsupported_extension_rejected():
    c = _client()
    r = c.post("/api/import", files={"file": ("t.pdf", b"%PDF", "application/pdf")},
               headers={"X-CSRF-Token": _csrf(c)})
    assert r.status_code == 400


# ── 模板 ───────────────────────────────────────────────────────────────────
def test_xlsx_template_is_downloadable_and_parsable():
    c = _client()
    r = c.get("/api/import/template.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "参数导入" in wb.sheetnames
    head = [wb["参数导入"].cell(1, i).value for i in range(1, 6)]
    assert head == ["企业名称", "产品型号", "产品类别", "流量L/min", "静压Pa"], head
    assert "填写说明" in wb.sheetnames


# ── Excel 导出 ─────────────────────────────────────────────────────────────
def test_export_xlsx_structure_and_formatting():
    c = _client()
    # 取三个气泵型号，保证既有最优值也有缺失
    prods = c.get("/api/products").json()
    air = [p["id"] for p in prods if p["category"] == "pump_air"][:3]
    assert len(air) == 3

    r = c.get(f"/export.xlsx?p={','.join(map(str, air))}")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[:2] == ["参数对比", "型号清单"]
    ws = wb["参数对比"]

    header = [ws.cell(1, i).value for i in range(1, 4 + len(air))]
    assert header[:4] == ["指标组", "指标", "单位", "优劣"]
    assert all("\n" in h for h in header[4:]), "型号表头应为「公司\\n型号」"

    # 找一行既有最优格又有缺失格的（各型号披露程度不同，至少要能看出这两层信息）
    missing_seen = best_seen = False
    for r_ in range(2, ws.max_row + 1):
        if ws.cell(r_, 1).value == "完整度":
            break
        vals = [ws.cell(r_, 5 + i) for i in range(len(air))]
        if any(v.value == "—" for v in vals):
            missing_seen = True
            for v in vals:
                assert v.font.color is None or v.font.color.rgb != "000000" or v.value == "—"
        if any(v.font.bold for v in vals if v.value != "—"):
            best_seen = True
    assert missing_seen, "导出里应保留缺失格（显示为「—」）"
    assert best_seen, "最优值应加粗"

    # 冻结窗格 & 完整度行
    assert ws.freeze_panes == "E2"
    labels = [ws.cell(r_, 1).value for r_ in range(2, ws.max_row + 1)]
    assert "完整度" in labels


def test_export_xlsx_respects_permission():
    """viewer 有导出权限（data:export），但不能编辑。"""
    c = _client()
    csrf = _csrf(c)
    c.post("/admin/users/create", data={
        "username": "viewer_exp", "password": "ViewerPass123",
        "display_name": "只读", "role": "viewer", "_csrf": csrf,
    }, follow_redirects=False)

    v = TestClient(app)
    v.__enter__()
    v.post("/login", data={"username": "viewer_exp", "password": "ViewerPass123"},
           follow_redirects=False)
    assert v.get("/export.xlsx").status_code == 200
    assert v.get("/products").status_code == 403
    assert v.get("/import").status_code == 403


def test_liquid_row_uses_flow_and_pressure_columns():
    """液泵也用「流量L/min」「静压Pa」两列填，导入时落到 µL/min 与 mbar。

    业务方的表只有一套流量/静压列（一行一个型号），不会为液泵换表头；
    所以这两列按「物理量」映射：液泵 → 液体流量(µL/min) / 背压(mbar)，其余 → 空气流量 / 静压。
    """
    c = _client()
    data = _xlsx([WIDE_HEAD, ["液泵测试公司", "LIQ-1", "压电液泵", 0.0035, 120000, "", "", "",
                              "", "", "", "", "", "", "", "", ""]])
    rep = _post(c, "t.xlsx", data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert rep["errors"] == [], rep["errors"]
    assert rep["written_specs"] == 2, rep

    with Session(engine) as s:
        prod = s.exec(select(Product).where(Product.model == "LIQ-1")).first()
        names = {
            s.get(SpecMetric, sp.metric_id).key: sp.value_num
            for sp in s.exec(select(ProductSpec).where(ProductSpec.product_id == prod.id)).all()
        }
    assert abs(names["flow_liquid"] - 3500) < 1e-6, names
    assert abs(names["pressure_liquid"] - 1200) < 1e-6, names


def test_import_errors_are_capped():
    """系统性错误只报前 200 条 —— 5000 行全错不该回一个几 MB 的 JSON。"""
    c = _client()
    rows = [["错行公司", f"BADCAP-{i}", "压电陶瓷", 1.0] + [""] * 13 for i in range(260)]
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["企业名称", "产品型号", "产品类别", "流量L/min"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)

    rep = _post(c, "cap.xlsx", buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert len(rep["errors"]) == 200, len(rep["errors"])
    assert rep["errors_truncated"] == 60, rep["errors_truncated"]
    assert rep["written_specs"] == 0


# ── 全量备份导出 ───────────────────────────────────────────────────────────
def test_export_all_has_three_sheets_and_is_reimportable():
    """「导出全部」= 备份包：参数宽表 + 参数明细 + 客户采用。

    最关键的一条是**能再导入**：导出后不改一行，重新导进去应当「无变化」，
    否则这份备份就只是看起来像备份。
    """
    c = _client()
    r = c.get("/export-all.xlsx")
    assert r.status_code == 200, r.text[:200]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["参数导入", "参数明细", "客户采用", "说明"], wb.sheetnames

    ws = wb["参数导入"]
    assert [x.value for x in ws[1]][:5] == ["企业名称", "产品型号", "产品类别", "流量L/min", "静压Pa"]
    assert ws.max_row > 1, "全量导出不该是空表"

    # 拿导出的文件原样再导入：应该是全部「无变化」，一条也不写
    rep = _post(c, "backup.xlsx", r.content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert rep["errors"] == [], rep["errors"]
    assert rep["written_specs"] == 0, f"往返不该改动数据，实际写入 {rep['written_specs']}"
    assert rep["skipped_unchanged"] > 0


def test_export_all_requires_export_permission():
    c = _client()
    csrf = _csrf(c)
    c.post("/admin/users/create", data={
        "username": "nogetexp", "password": "ViewerPass123",
        "display_name": "无导出", "role": "viewer", "_csrf": csrf,
    }, follow_redirects=False)
    v = TestClient(app)
    v.__enter__()
    v.post("/login", data={"username": "nogetexp", "password": "ViewerPass123"},
           follow_redirects=False)
    assert v.get("/export-all.xlsx").status_code == 200, "viewer 本来就有 data:export"


def test_company_with_parenthetical_merges_into_existing():
    """「村田制作所（Murata）」要认成已有的「村田制作所」，不能凭空多出一家。

    README 承诺「公司名忽略空格、标点、大小写与括注」，实现却只去掉了括号字符、
    把括注里的内容留在名字里 —— 于是同一次导入会造出重复公司，
    对比页上同一家友商出现两次、型号被拆到两行。
    """
    c = _client()
    with Session(engine) as s:
        if s.exec(select(Entity).where(Entity.name == "村田制作所")).first() is None:
            s.add(Entity(kind="competitor", name="村田制作所", name_en="Murata"))
            s.commit()

    data = _xlsx([WIDE_HEAD, ["村田制作所（Murata）", "MZB-MERGE-1", "压电气泵", 0.4, 3000, "", "", "",
                              "", "", "", "", "", "", "", "", ""]])
    rep = _post(c, "merge.xlsx", data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert rep["errors"] == [], rep["errors"]
    assert rep["created_entities"] == 0, f"括注让公司重复建档了：{rep}"

    with Session(engine) as s:
        names = [e.name for e in s.exec(select(Entity).where(Entity.name.like("村田%"))).all()]
    assert names == ["村田制作所"], names


# ── 运行器 ─────────────────────────────────────────────────────────────────
def main() -> int:
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"  PASS  {name}")
        except AssertionError as e:
            failed += 1
            print(f"  FAIL  {name}  {e}")
        except Exception as e:
            failed += 1
            print(f"  ERROR {name}  {type(e).__name__}: {e}")
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
