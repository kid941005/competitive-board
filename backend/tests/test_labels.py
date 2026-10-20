"""界面文案回归：自家公司直接用公司名，不出现「本方 / 我方 / 本公司」代称。

这条要求很具体也容易被后续改版改回去，所以用断言锁住：
  · 对比页的型号选择器里，友商带「友商」标签，自家公司裸名称
  · 「只看 ××」按钮用的是公司名
  · 导入页、产品管理页的公司下拉不带「（本方）」后缀
  · 全文不出现「本方 / 我方 / 本公司」

两种跑法都支持：pytest tests/ 或 python tests/test_labels.py
"""

from __future__ import annotations

import os
import re
import sys
import tempfile
from contextlib import contextmanager
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="cib-label-test-")
os.environ["CIB_DATA_DIR"] = _TMP
os.environ["CIB_ADMIN_USER"] = "admin"
os.environ["CIB_ADMIN_PASSWORD"] = "AdminPass123"
os.environ.pop("DATABASE_URL", None)

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient  # noqa: E402
from sqlmodel import Session, select  # noqa: E402

from app.db import engine, init_db  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Entity, Product  # noqa: E402
from app.seed import seed  # noqa: E402

ADMIN_PW = "AdminPass123"
SELF_NAME = "左蓝微"
BANNED = ("本方", "我方", "本公司", "（本方）")

_client: TestClient | None = None


def _c() -> TestClient:
    global _client
    if _client is not None:
        return _client
    init_db()
    seed()
    c = TestClient(app)
    c.__enter__()
    c.post("/login", data={"username": "admin", "password": ADMIN_PW}, follow_redirects=False)
    _client = c
    return c


@contextmanager
def _self_product():
    """临时给自家公司挂一个型号 —— 选择器只渲染有型号的公司。

    自家型号要等用户录入，测试里不能因为「还没有」就把渲染逻辑漏测。
    """
    with Session(engine) as s:
        ent = s.exec(select(Entity).where(Entity.kind == "self")).first()
        existing = s.exec(select(Product).where(Product.entity_id == ent.id)).first()
    temp_id = None
    if existing is None:
        with Session(engine) as s:
            prod = Product(entity_id=ent.id, model="SELF-TMP", name="自家临时型号", category="pump_air")
            s.add(prod)
            s.commit()
            s.refresh(prod)
            temp_id = prod.id
    try:
        yield
    finally:
        if temp_id is not None:
            with Session(engine) as s:
                obj = s.get(Product, temp_id)
                if obj is not None:
                    s.delete(obj)
                    s.commit()


def _assert_clean(html: str, page: str) -> None:
    for word in BANNED:
        assert word not in html, f"{page} 仍出现「{word}」"


# ── 数据层 ─────────────────────────────────────────────────────────────────
def test_self_entity_uses_real_name():
    """kind=self 的实体名就是公司名本身，不是「本公司」这种占位。"""
    _c()
    with Session(engine) as s:
        ent = s.exec(select(Entity).where(Entity.kind == "self")).first()
    assert ent is not None, "缺少 kind=self 的实体"
    assert ent.name == SELF_NAME, f"期望 {SELF_NAME}，实际 {ent.name}"


def test_legacy_self_name_is_migrated():
    """老数据卷里叫「本公司」的，启动时会被改成真实公司名。

    这条针对的是升级场景：代码更新了但数据库还是旧的。
    """
    from app.main import _rename_legacy_self

    with Session(engine) as s:
        ent = s.exec(select(Entity).where(Entity.kind == "self")).first()
        ent.name = "本公司"
        s.add(ent)
        s.commit()

    with Session(engine) as s:
        _rename_legacy_self(s)

    with Session(engine) as s:
        ent = s.exec(select(Entity).where(Entity.kind == "self")).first()
        assert ent.name == SELF_NAME, f"应迁移为 {SELF_NAME}，实际 {ent.name}"


# ── 对比页 ─────────────────────────────────────────────────────────────────
def test_compare_picker_shows_plain_name():
    """自家公司在选择器里只显示名字；友商才带「友商」标签。

    选择器只渲染「有型号」的公司，而自家型号要等你录入才知道，
    所以这里临时建一个自家型号来验证渲染，用完删掉。
    """
    c = _c()
    with _self_product():
        html = c.get("/compare").text
        _assert_clean(html, "/compare")

        # 自家那一行：有 data-kind="self"，标签块里不应有类别徽标
        row = re.search(
            r'<div class="ent-row" data-kind="self">(.*?)<div class="chips">',
            html,
            re.S,
        )
        assert row is not None, "对比页没有渲染自家公司行"
        assert SELF_NAME in row.group(1), "自家公司名未显示"
        assert 'class="ent-kind' not in row.group(1), "自家公司不应带类别徽标"

        # 友商仍要带徽标，否则「谁是友商」这个信息就丢了
        assert 'class="ent-kind competitor">友商' in html


def test_only_self_button_uses_company_name():
    """「只看 ××」按钮用公司名，不写「只看本方」。"""
    c = _c()
    with _self_product():
        html = c.get("/compare").text
        assert f"只看{SELF_NAME}" in html, "按钮文案应为「只看公司名」"
        assert "只看本方" not in html


# ── 导入页 / 产品管理页 ────────────────────────────────────────────────────
def test_company_dropdowns_have_no_suffix():
    c = _c()
    for path in ("/datasheet", "/products"):
        html = c.get(path).text
        _assert_clean(html, path)
        assert SELF_NAME in html, f"{path} 的公司下拉里应有 {SELF_NAME}"


def test_export_category_label_for_self():
    """导出「类别」列：自家记为「自研」，友商仍是「友商」。"""
    c = _c()
    with Session(engine) as s:
        ent = s.exec(select(Entity).where(Entity.kind == "self")).first()
        pid = s.exec(select(Product).where(Product.entity_id == ent.id)).first()
    if pid is None:
        return  # 自家还没有型号，跳过

    import io

    from openpyxl import load_workbook

    r = c.get(f"/export.xlsx?p={pid.id}")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content))["型号清单"]
    kinds = {ws.cell(i, 1).value: ws.cell(i, 4).value for i in range(2, ws.max_row + 1)}
    assert kinds.get(SELF_NAME) == "自研", kinds


# ── 运行器 ─────────────────────────────────────────────────────────────────
def main() -> int:
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"  PASS  {name}")
        except AssertionError as e:
            failed += 1
            print(f"  FAIL  {name}  {e}")
        except Exception as e:  # noqa: BLE001
            failed += 1
            print(f"  ERROR {name}  {type(e).__name__}: {e}")
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
